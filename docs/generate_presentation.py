# ============================================================
# docs/generate_presentation.py
# SMART CROP ROTATION - generates the PowerPoint presentation.
# ------------------------------------------------------------
# Optional dependency:
#     pip install python-pptx
# Run (from the project root):
#     venv\\Scripts\\python.exe docs\\generate_presentation.py
# Output: docs/smart_crop_rotation_presentation.pptx
# ============================================================

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

GREEN = RGBColor(0x2E, 0x7D, 0x32)
DARK_GREEN = RGBColor(0x1B, 0x5E, 0x20)
GOLD = RGBColor(0xF9, 0xA8, 0x25)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x45, 0x5A, 0x64)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "smart_crop_rotation_presentation.pptx")

SLIDES = [
    ("SMART CROP ROTATION",
     ["AI-Based Intelligent Crop Rotation Recommendation System",
      "BCA Final-Year Project", "Your Name | Your College"]),
    ("Introduction",
     ["Farmers must decide: what to grow in the next season?",
      "The choice depends on soil, season, water, weather and previous crop.",
      "This project recommends the best next crop using ML + agriculture rules."]),
    ("Problem Statement",
     ["Repeating the same crop degrades soil and increases disease risk.",
      "Existing tools recommend only one crop - without explanation.",
      "They ignore previous crop, crop families, disease and fertiliser history."]),
    ("Existing System",
     ["Experience-based decisions and generic crop calendars.",
      "Simple apps using only 2-3 inputs (NPK, temperature).",
      "No explanation, no history, no rotation logic."]),
    ("Proposed System",
     ["Web application with 17 farm inputs.",
      "ML prediction combined with a rule-based rotation engine.",
      "Top crop + 3 alternatives, fully explained."]),
    ("Objectives",
     ["Train and compare 4 Machine Learning models.",
      "Build a crop rotation engine from agricultural rules.",
      "Compute Soil Health, Rotation and Sustainability scores.",
      "Explanations, history, analytics and printable reports."]),
    ("Technologies Used",
     ["Frontend: HTML, CSS, JavaScript, Chart.js",
      "Backend: Python + Flask",
      "Data: Pandas, NumPy",
      "Machine Learning: Scikit-learn, Joblib (model storage)",
      "Database: SQLite"]),
    ("System Architecture",
     ["Browser -> Flask application",
      "-> Rotation engine + Soil analysis + SQLite + ML model",
      "-> Explained recommendation result"]),
    ("Modules",
     ["Data generation, Model training, Rotation engine",
      "Soil health analysis, Web application",
      "History, Analytics, Reports, REST API"]),
    ("Dataset",
     ["~700 synthetic records, 17 crops, 16 features.",
      "Generated with realistic agronomic relationships - not random noise.",
      "Clearly labelled as demo data."]),
    ("Machine Learning",
     ["Algorithms: Decision Tree, Random Forest, Gradient Boosting, KNN.",
      "80/20 train-test split (stratified).",
      "Metrics: Accuracy, Precision, Recall, F1, Confusion Matrix.",
      "Best model auto-selected and saved with Joblib."]),
    ("Crop Rotation Algorithm",
     ["Final score = 45% ML + 40% rotation rules + 15% environment.",
      "Rotation score = family 25 + soil 25 + water 20 + disease 15 + season 15."]),
    ("Soil Health Analysis",
     ["Soil Health Score 0-100: N, P, K, pH, fertiliser, previous crop.",
      "Sustainability Score 0-100: water, fertiliser, diversity, soil, disease."]),
    ("UI Screenshots",
     ["Home dashboard with score cards.",
      "Recommendation form and explained result page.",
      "History, analytics charts, model performance page."]),
    ("Recommendation Result",
     ["Example: Previous Crop = Rice  ->  Recommended = Green Gram",
      "Alternatives: Maize, Groundnut, Black Gram.",
      "Each with Suitability %, Rotation, Soil, Water and Season info.",
      "\"Why this crop?\" checklist and nitrogen advice."]),
    ("Model Performance",
     ["Table + chart comparing the 4 classifiers.",
      "Confusion matrix for the selected model.",
      "Metrics labelled as trained on synthetic demo data."]),
    ("Advantages",
     ["Explainable recommendations (Why this crop?).",
      "Lightweight, local, easy to demonstrate.",
      "Combines ML, rules, soil analysis and sustainability.",
      "History, analytics and REST API included."]),
    ("Future Enhancement",
     ["Live weather and soil-test APIs, real datasets.",
      "Market-price optimisation and deep learning.",
      "Mobile app and multi-language support."]),
    ("Conclusion",
     ["Successfully integrates web development + database + machine learning.",
      "A practical, honest decision-support system - not a guaranteed-yield tool."]),
    ("Thank You",
     ["Questions?"]),
]


def add_title_slide(prs, title, points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK_GREEN

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(2.4), Inches(9.2), Inches(1.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(8.8), Inches(2.4))
    tf2 = tb2.text_frame
    for i, line in enumerate(points):
        para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        para.text = line
        para.font.size = Pt(20 if i == 0 else 16)
        para.font.color.rgb = GOLD if i == 0 else WHITE
        para.alignment = PP_ALIGN.CENTER


def add_content_slide(prs, title, points):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE

    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9.0), Inches(0.9))
    title_p = box.text_frame.paragraphs[0]
    title_p.text = title
    title_p.font.size = Pt(32)
    title_p.font.bold = True
    title_p.font.color.rgb = GREEN

    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(9.0), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()

    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.6), Inches(5.2))
    tf = body.text_frame
    tf.word_wrap = True
    for i, pt in enumerate(points):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = ("\u2022  " if pt != "Questions?" else "? ") + pt
        para.font.size = Pt(20)
        para.font.color.rgb = GREY
        para.space_after = Pt(12)
        para.level = 1 if "\u2022  " not in para.text and pt != "Questions?" else 0


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for i, (title, points) in enumerate(SLIDES):
        if i == 0 or title == "Thank You":
            add_title_slide(prs, title, points)
        else:
            add_content_slide(prs, title, points)

    prs.save(OUT)
    print(f"[OK] Presentation saved to {OUT}")


if __name__ == "__main__":
    main()